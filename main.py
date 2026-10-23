import os
import tempfile
import math
import shutil
from fastapi import FastAPI, HTTPException
from fastapi.responses import JSONResponse
from pydantic import BaseModel
from supabase import create_client, Client
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import requests

app = FastAPI(title="VPP Report API", version="1.0")

# Credenciales de Supabase (reemplaza con tus variables de entorno en producción)
SUPABASE_URL = "https://slylehtfcliysnodbidy.supabase.co"
SUPABASE_KEY = "tu_service_role_key_o_anon_key"
supabase: Client = create_client(SUPABASE_URL, SUPABASE_KEY)

class ReportRequest(BaseModel):
    codigo_reporte: str

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_cell_margins(cell, top=60, bottom=60, left=60, right=60):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def aplicar_fuente_poppins_9(run, bold=False, color_rgb=(30, 41, 59)):
    run.font.name = "Poppins"
    run.font.size = Pt(9)
    run.bold = bold
    run.italic = False
    run.font.color.rgb = RGBColor(*color_rgb)

def procesar_parrafo(p, metadata):
    for run in p.runs:
        for k, v in metadata.items():
            clave = f"{{{{{k}}}}}"
            if clave in run.text:
                run.text = run.text.replace(clave, str(v))
                aplicar_fuente_poppins_9(run, bold=False)
                
    full_text = p.text
    cambiado = False
    for k, v in metadata.items():
        clave = f"{{{{{k}}}}}"
        if clave in full_text:
            full_text = full_text.replace(clave, str(v))
            cambiado = True
            
    if cambiado:
        p.text = full_text
        for run in p.runs:
            aplicar_fuente_poppins_9(run, bold=False)

def procesar_contenedor(contenedor, metadata):
    for p in contenedor.paragraphs:
        procesar_parrafo(p, metadata)
    if hasattr(contenedor, "tables"):
        for tabla in contenedor.tables:
            for fila in tabla.rows:
                for celda in fila.cells:
                    procesar_contenedor(celda, metadata)

def rellenar_tabla_dinamica(tabla, filas_datos, mapeo_columnas):
    if len(tabla.rows) < 2:
        return
    fila_plantilla = tabla.rows[1]
    if not filas_datos:
        return

    for idx, item in enumerate(filas_datos):
        fila_actual = tabla.add_row() if idx > 0 else fila_plantilla
        color_fondo = "FFFFFF" if idx % 2 == 0 else "F8FAFC"
        
        for c_idx, cell in enumerate(fila_actual.cells):
            set_cell_background(cell, color_fondo)
            set_cell_margins(cell, top=50, bottom=50, left=60, right=60)
            
            if c_idx in mapeo_columnas:
                key = mapeo_columnas[c_idx]
                if callable(key):
                    val = key(idx, item)
                else:
                    val = item.get(key, "")
                cell.text = str(val)
                for p in cell.paragraphs:
                    for run in p.runs:
                        aplicar_fuente_poppins_9(run, bold=False)

@app.post("/api/generar-reporte")
def generar_reporte(payload: ReportRequest):
    codigo = payload.codigo_reporte
    
    # 1. Consultar datos principales en Supabase
    res_repo = supabase.from_("reportes_diarios").select("*").eq("codigo_reporte", codigo).execute()
    if not res_repo.data:
        raise HTTPException(status_code=404, detail="Reporte no encontrado en Supabase")
    
    repo_data = res_repo.data[0]
    
    # 2. Consultar tablas de detalle
    dotacion = supabase.from_("reporte_dotacion").select("*").eq("codigo_reporte", codigo).execute().data or []
    equipos = supabase.from_("reporte_equipos").select("*").eq("codigo_reporte", codigo).execute().data or []
    actividades = supabase.from_("reporte_actividades").select("*").eq("codigo_reporte", codigo).execute().data or []
    fotos = supabase.from_("reporte_fotos").select("*").eq("codigo_reporte", codigo).execute().data or []

    # 3. Preparar metadatos para la plantilla
    metadata = {
        "id": repo_data.get("codigo_reporte", codigo),
        "wbs": "3400",
        "contratista": repo_data.get("contratista", ""),
        "emisor": repo_data.get("emisor_nombre", ""),
        "aprobador": repo_data.get("aprobador_asignado", ""),
        "proyecto": "1008765 24 PULG.STP PHASE 2A",
        "id_contrato": repo_data.get("contrato", ""),
        "fecha": str(repo_data.get("fecha", "")),
        "hora_ingreso_area": "08:00",
        "hora_salida_almuerzo": "13:00",
        "hora_entrada_almuerzo": "14:00",
        "hora_salida_area": "20:00",
        "emisor_firma": f"{repo_data.get('emisor_nombre', '')} - Inspector Técnico",
        "aprobador_firma": f"{repo_data.get('aprobador_asignado', '')} - Aprobador VPP"
    }

    # 4. Cargar plantilla local
    ruta_plantilla = "Reporte Diario.docx"
    if not os.path.exists(ruta_plantilla):
        raise HTTPException(status_code=500, detail="Plantilla Word no encontrada en el servidor")

    doc = Document(ruta_plantilla)

    # Procesar contenedores estándar y XML crudo
    containers = [doc]
    for section in doc.sections:
        for h in [section.header, section.first_page_header, section.even_page_header]:
            if h: containers.append(h)
        for f in [section.footer, section.first_page_footer, section.even_page_footer]:
            if f: containers.append(f)

    for cont in containers:
        procesar_contenedor(cont, metadata)

    for part in doc.part.package.parts:
        if 'xml' in part.content_type:
            try:
                xml_str = part.blob.decode('utf-8')
                modified = False
                for k, v in metadata.items():
                    clave = f"{{{{{k}}}}}"
                    if clave in xml_str:
                        xml_str = xml_str.replace(clave, str(v))
                        modified = True
                if modified:
                    part._blob = xml_str.encode('utf-8')
            except:
                pass

    # Guardar temporalmente
    temp_dir = tempfile.mkdtemp()
    docx_path = os.path.join(temp_dir, f"{codigo}.docx")
    doc.save(docx_path)

    # 5. Rellenar tablas dinámicas
    doc_final = Document(docx_path)
    for tabla in doc_final.tables:
        if len(tabla.rows) > 0:
            texto_header = "".join([c.text.upper() for c in tabla.rows[0].cells])
            if "ACTIVIDAD" in texto_header and "AVANCE FÍSICO" in texto_header:
                while len(tabla.rows) > 2:
                    tabla._tbl.remove(tabla.rows[2]._tr)
                rellenar_tabla_dinamica(tabla, actividades, {0: lambda i, x: str(i + 1), 1: "actividad_nombre", 2: lambda i, x: f"{x.get('avance_diario', 0)} {x.get('unidad', '')}"})
            elif "CARGO" in texto_header and "HH" in texto_header:
                while len(tabla.rows) > 2:
                    tabla._tbl.remove(tabla.rows[2]._tr)
                rellenar_tabla_dinamica(tabla, dotacion, {0: lambda i, x: str(i + 1), 1: "cargo_nombre", 2: "cantidad", 3: lambda i, x: f"{x.get('hh_total', 0)} HH"})
            elif "EQUIPO" in texto_header and "HM" in texto_header:
                while len(tabla.rows) > 2:
                    tabla._tbl.remove(tabla.rows[2]._tr)
                rellenar_tabla_dinamica(tabla, equipos, {0: lambda i, x: str(i + 1), 1: "equipo_nombre", 2: "cantidad", 3: lambda i, x: f"{x.get('hm_total', 0)} HM"})

    # 6. Inyectar Fotos en Grilla de 2 Columnas
    for p in doc_final.paragraphs:
        if "{{REGISTRO_FOTOGRAFICO}}" in p.text:
            p.text = ""
            tabla_fotos = doc_final.add_table(rows=0, cols=2)
            tabla_fotos.autofit = False
            p._p.addprevious(tabla_fotos._tbl)
            
            num_filas = math.ceil(len(fotos) / 2.0)
            for r_idx in range(num_filas):
                row = tabla_fotos.add_row()
                for c_idx in range(2):
                    foto_idx = (r_idx * 2) + c_idx
                    if foto_idx < len(fotos):
                        foto_item = fotos[foto_idx]
                        cell = row.cells[c_idx]
                        set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
                        
                        # Descargar imagen desde Supabase Storage URL
                        img_url = foto_item.get("url_foto")
                        local_img = os.path.join(temp_dir, f"foto_{foto_idx}.jpg")
                        if img_url:
                            try:
                                r = requests.get(img_url)
                                with open(local_img, "wb") as f:
                                    f.write(r.content)
                            except:
                                pass

                        p_img = cell.paragraphs[0]
                        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        run_img = p_img.add_run()
                        if os.path.exists(local_img):
                            run_img.add_picture(local_img, width=Inches(3.0))
                        
                        p_desc = cell.add_paragraph()
                        p_desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        run_pref = p_desc.add_run(f"Foto {foto_idx + 1} - ")
                        aplicar_fuente_poppins_9(run_pref, bold=True, color_rgb=(0, 32, 96))
                        run_txt = p_desc.add_run(foto_item.get("descripcion", "Inspección visual en terreno."))
                        aplicar_fuente_poppins_9(run_txt, bold=False, color_rgb=(30, 41, 59))
            break

    doc_final.save(docx_path)
    
    # Subir archivo generado a Supabase Storage
    with open(docx_path, "rb") as f:
        file_bytes = f.read()
    
    storage_path = f"generados/{codigo}.docx"
    supabase.storage.from_("reportes_generados").upload(storage_path, file_bytes, {"file-option": "upsert"})
    public_url = supabase.storage.from_("reportes_generados").get_public_url(storage_path)

    return JSONResponse(content={
        "status": "success",
        "codigo_reporte": codigo,
        "url_documento": public_url
    })